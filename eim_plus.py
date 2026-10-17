"""
EIM+  -  the engine layer on top of app.py (app.py and train_grpo.py stay untouched)
=====================================================================================

What it adds, mapped to the known limitations of the base loop:

    limitation (base EIM)                    -> what EIM+ does
    ---------------------------------------- ---------------------------------------------------------------
    overfitting to the visible tests         -> hold-out tests: the loop selects on a visible split, the final
                                                answer is re-checked on tests it never saw; a hard-coding detector
                                                flags solutions that just embed the expected values
    shallow fuzzy memory (difflib only)      -> SmartMemory: same-failure-class gate + char n-gram cosine + difflib
                                                + exception-class bonus, similarity-weighted voting, recency boost
    local minima / early stall               -> escalation: after 2 stale rounds the loop RESTARTS from the task
                                                (fresh, high-temperature candidates, "use a different algorithm")
    code can only be generated from scratch  -> the loop can start from YOUR code (fix / optimise an uploaded file)
    tests had to be typed by hand            -> tests are extracted from asserts, fenced blocks, or pytest-style files
    answer is only as good as the tests      -> + test linting (contradictions, trivial asserts) and MUTATION CHECK: small
                                                bugs are injected into the final code; if the tests still pass, they are
                                                too weak and the report says so
    CPU/GPU cost and latency                 -> hardware profile (lean / balanced / full) picks safe defaults; candidates
                                                are generated in a ladder (1 first, the rest only if it did not help);
                                                parsed/BM25 indexes and memory n-grams are cached; optional time budget;
                                                successful runs are replayed from a cache
    sandbox: no network, small memory        -> pre-flight: missing/heavy/network modules are detected BEFORE spending
                                                iterations; the model is told the network is off; unrunnable tests stop
                                                the loop instead of burning attempts
    learns only a generic strategy           -> RepairLog: verified (task, broken code, failing tests, fixed code) records,
                                                retrieved by similarity and shown to the model as worked examples; weak-test
                                                or hard-coded runs are never learned from; --export-sft makes LoRA/QLoRA data
    free Hugging Face ZeroGPU Space          -> detected (SPACES_ZERO_GPU): the loop stops starting new rounds at 65% of the
                                                GPU call's duration (EIM_GPU_SECONDS) and returns its best version instead
                                                of being killed with nothing; optional EIM_MEMORY_REPO keeps the learned
                                                memory across the Space's restarts (its disk is wiped)
    max_iter spent without a full solution   -> adaptive budget: extra rounds are granted while the code is still
                                                improving (bounded by the profile), the best version is always kept

Plus the "read anything" layer for the chat app: text/code/pdf/docx/xlsx/csv/ipynb/zip/images, and a dependency-free
BM25 retriever (Arabic-aware normalisation) that picks the relevant parts of big files for the model's small context.

Self-test (no model, no network):   python eim_plus.py --selftest
"""
from __future__ import annotations

import ast
import copy
import difflib
import functools
import hashlib
import importlib.util
import json
import math
import os
import platform
import re
import shutil
import sys
import tempfile
import threading
import time
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from typing import Iterator, Sequence

from app import (CFG, DCME, EUTV, IFC, STRATEGIES, Config, ExecResult, ExperienceMemory, LanguageModel, Mutation,
                 Step, _gpu_duration_seconds, _is_zero_gpu_space, _repair_rules, draft_prompt, extract_code, normalise, score)


# ============================================================================
# 0. Hardware profile: safe defaults for weak machines, full power for strong ones
# ============================================================================
@dataclass(frozen=True)
class Profile:
    name: str            # lean | balanced | full
    cpus: int
    ram_gb: float        # 0.0 = unknown
    cuda: bool
    k_cap: int           # most candidates per round this machine should generate
    ladder: bool         # generate 1 candidate first, the rest only if needed
    mutants: int         # mutants used by the mutation check (0 = off)
    bonus_iters: int     # extra rounds granted while the code is still improving
    zero_gpu: bool = False   # Hugging Face ZeroGPU: GPU time is metered per call, the disk is wiped on restart


_PROFILES = {
    "lean": dict(k_cap=3, ladder=True, mutants=3, bonus_iters=1),
    "balanced": dict(k_cap=4, ladder=True, mutants=5, bonus_iters=2),
    "full": dict(k_cap=8, ladder=False, mutants=8, bonus_iters=2),
}


def _ram_gb() -> float:
    """TOTAL RAM in GB (0.0 if it cannot be determined). Total, not free: the profile describes the machine, and free
    memory drops as soon as the model is loaded. No hard dependency on psutil."""
    try:
        import psutil
        return psutil.virtual_memory().total / 2 ** 30
    except Exception:
        pass
    try:
        with open("/proc/meminfo", encoding="ascii", errors="ignore") as handle:
            for line in handle:
                if line.startswith("MemTotal:"):
                    return int(line.split()[1]) / 2 ** 20
    except Exception:
        pass
    try:
        return os.sysconf("SC_PAGE_SIZE") * os.sysconf("SC_PHYS_PAGES") / 2 ** 30
    except Exception:
        return 0.0


def _has_accelerator() -> bool:
    """CUDA / Apple-silicon / ZeroGPU. Works before torch or model weights load."""
    if _is_zero_gpu_space():
        return True
    torch = sys.modules.get("torch")
    if torch is not None:
        try:
            mps = getattr(getattr(torch, "backends", None), "mps", None)
            return bool(torch.cuda.is_available() or (mps is not None and mps.is_available()))
        except Exception:
            return False
    return (bool(shutil.which("nvidia-smi")) or os.path.exists("/proc/driver/nvidia/version")
            or (platform.system() == "Darwin" and platform.machine() == "arm64"))


@functools.lru_cache(maxsize=1)
def hardware_profile() -> Profile:
    """EIM_PROFILE=lean|balanced|full forces a profile; otherwise it is derived from CPU, RAM and CUDA."""
    cpus, ram, cuda = os.cpu_count() or 1, _ram_gb(), _has_accelerator()
    forced = os.environ.get("EIM_PROFILE", "").strip().lower()
    if forced in _PROFILES:
        name = forced
    elif cuda or (cpus >= 8 and ram >= 16):
        name = "full"
    elif cpus <= 4 or (0 < ram < 8):
        name = "lean"
    else:
        name = "balanced"
    params = dict(_PROFILES[name])
    zero = _is_zero_gpu_space()
    if zero:                      # every second inside the GPU call is metered: fewer mutants, at most one bonus round
        params.update(mutants=min(params["mutants"], 5), bonus_iters=min(params["bonus_iters"], 1))
    return Profile(name, cpus, round(ram, 1), cuda, zero_gpu=zero, **params)


def _env_float(name: str, default: float) -> float:
    try:
        return float(os.environ.get(name, default))
    except ValueError:
        return default


def effective_time_budget(profile: Profile, explicit: float | None = None) -> float:
    """Seconds after which no NEW round starts (0 = unlimited). Priority: argument, then EIM_TIME_BUDGET, then
    on ZeroGPU only 65% of the actual clamped @spaces.GPU duration (120 s for large, 60 s for xlarge by default).
    This returns the best version found before the platform ends the GPU call."""
    if explicit is not None:
        return float(explicit)
    configured = _env_float("EIM_TIME_BUDGET", 0.0)
    if configured > 0 or not profile.zero_gpu:
        return configured
    return 0.65 * _gpu_duration_seconds()


# ============================================================================
# 1. Attachments: read anything into text
# ============================================================================
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_PDF_PAGES = 300
MAX_TABLE_ROWS = 300
MAX_ZIP_MEMBERS = 30
MAX_ZIP_MEMBER_BYTES = 400_000

CODE_EXT = {".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".rb",
            ".php", ".sh", ".sql", ".kt", ".swift", ".lua", ".r", ".css", ".html", ".xml", ".yaml", ".yml", ".toml",
            ".ini", ".cfg", ".json", ".jsonl"}
TEXT_EXT = {".txt", ".md", ".rst", ".log", ".tex"}
TEXT_NAMES = {"dockerfile", "makefile", "readme", "license", "requirements.txt"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}


@dataclass
class Attachment:
    name: str
    path: str
    kind: str                  # code | text | document | table | image | archive | unsupported
    text: str = ""
    note: str = ""
    size: int = 0
    image_size: tuple[int, int] | None = None


def decode_bytes(data: bytes) -> str:
    """UTF-8 first, then UTF-16 (BOM), then cp1256 (Windows Arabic), finally latin-1 (never fails)."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    for encoding in ("utf-8-sig", "cp1256"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _is_binary(data: bytes) -> bool:
    return b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff"))


def _read_pdf(path: str) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0
        if not unlocked:
            return "", "the PDF is password protected"
    pages = []
    for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
        content = (page.extract_text() or "").strip()
        if content:
            pages.append(f"[page {number}]\n{content}")
    note = ""
    if len(reader.pages) > MAX_PDF_PAGES:
        note = f"only the first {MAX_PDF_PAGES} of {len(reader.pages)} pages were read"
    if not pages:
        note = "no extractable text (a scanned PDF needs OCR)"
    return "\n\n".join(pages), note


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _safe_xml(data: bytes) -> ET.Element:
    if b"<!DOCTYPE" in data or b"<!ENTITY" in data:   # entity-expansion attacks
        raise ValueError("DTD/entities are not allowed")
    return ET.fromstring(data)


def _read_docx(path: str) -> tuple[str, str]:
    with zipfile.ZipFile(path) as archive:
        info = archive.getinfo("word/document.xml")
        if info.file_size > 60 * 1024 * 1024:
            return "", "document.xml is too large"
        root = _safe_xml(archive.read("word/document.xml"))
    paragraphs = []
    for paragraph in root.iter(f"{_W}p"):
        text = "".join(node.text or "" for node in paragraph.iter(f"{_W}t")).strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs), ""


def _read_xlsx(path: str) -> tuple[str, str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    parts, note = [], ""
    for sheet in workbook.worksheets[:10]:
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= MAX_TABLE_ROWS:
                note = f"sheets truncated to {MAX_TABLE_ROWS} rows"
                break
            rows.append(",".join("" if cell is None else str(cell) for cell in row))
        parts.append(f"[sheet {sheet.title}]\n" + "\n".join(rows))
    workbook.close()
    return "\n\n".join(parts), note


def _read_ipynb(text: str) -> str:
    notebook = json.loads(text)
    out = []
    for cell in notebook.get("cells", []):
        source = cell.get("source", "")
        source = "".join(source) if isinstance(source, list) else str(source)
        if cell.get("cell_type") == "code":
            out.append(f"```python\n{source}\n```")
        else:
            out.append(source)
    return "\n\n".join(out)


def _read_zip(path: str) -> tuple[str, str]:
    parts, note = [], ""
    with zipfile.ZipFile(path) as archive:
        members = [m for m in archive.infolist() if not m.is_dir()]
        listing = "\n".join(f"- {m.filename} ({m.file_size} bytes)" for m in members[:200])
        parts.append(f"[archive listing]\n{listing}")
        read = 0
        for member in members:
            ext = os.path.splitext(member.filename)[1].lower()
            if ext not in CODE_EXT | TEXT_EXT or member.file_size > MAX_ZIP_MEMBER_BYTES:
                continue
            if read >= MAX_ZIP_MEMBERS:
                note = f"only the first {MAX_ZIP_MEMBERS} text files of the archive were read"
                break
            data = archive.read(member)       # never extracted to disk: path tricks are irrelevant
            if not _is_binary(data):
                parts.append(f"[file {member.filename}]\n{decode_bytes(data)}")
                read += 1
    return "\n\n".join(parts), note


def read_attachment(path: str) -> Attachment:
    name = os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()
    try:
        size = os.path.getsize(path)
    except OSError as exc:
        return Attachment(name, path, "unsupported", note=f"cannot open: {exc}")
    att = Attachment(name, path, "unsupported", size=size)
    if size > MAX_FILE_BYTES:
        att.note = f"file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB"
        return att
    try:
        if ext in IMAGE_EXT:
            from PIL import Image

            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image:
                att.image_size = image.size
            att.kind = "image"
        elif ext == ".pdf":
            att.text, att.note = _read_pdf(path)
            att.kind = "document"
        elif ext == ".docx":
            att.text, att.note = _read_docx(path)
            att.kind = "document"
        elif ext in {".xlsx", ".xlsm"}:
            att.text, att.note = _read_xlsx(path)
            att.kind = "table"
        elif ext == ".zip":
            att.text, att.note = _read_zip(path)
            att.kind = "archive"
        elif ext in CODE_EXT | TEXT_EXT or ext in {".csv", ".tsv", ".ipynb"} or name.lower() in TEXT_NAMES:
            with open(path, "rb") as handle:
                data = handle.read(MAX_FILE_BYTES)
            if _is_binary(data):
                att.note = "binary file"
                return att
            text = decode_bytes(data)
            if ext == ".ipynb":
                att.text, att.kind = _read_ipynb(text), "code"
            elif ext in {".csv", ".tsv"}:
                lines = text.splitlines()
                if len(lines) > MAX_TABLE_ROWS:
                    att.note = f"{len(lines)} rows, showing the first {MAX_TABLE_ROWS}"
                att.text, att.kind = "\n".join(lines[:MAX_TABLE_ROWS]), "table"
            else:
                att.text, att.kind = text, ("code" if ext in CODE_EXT else "text")
        else:
            att.note = f"unsupported file type {ext or '(none)'}"
    except Exception as exc:    # corrupt file, missing optional dependency, ...
        att.kind, att.text = "unsupported", ""
        att.note = f"could not read ({type(exc).__name__}: {exc})"
    return att


# ============================================================================
# 2. Retrieval: BM25 over chunks (Arabic-aware, no dependencies)
# ============================================================================
_AR_MARKS = re.compile("[\u064B-\u0652\u0640]")          # tashkeel + tatweel
_AR_MAP = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه"})


def tokenize(text: str) -> list[str]:
    text = _AR_MARKS.sub("", text).translate(_AR_MAP).lower()
    tokens = re.findall(r"\w+", text)
    return tokens + [part for t in tokens if "_" in t for part in t.split("_") if part]


def chunk_text(text: str, size: int = 1200, overlap: int = 150) -> list[str]:
    pieces: list[str] = []
    for line in text.splitlines(keepends=True):
        while len(line) > size:
            pieces.append(line[:size])
            line = line[size:]
        pieces.append(line)
    chunks: list[str] = []
    current: list[str] = []
    length = 0
    for piece in pieces:
        if length + len(piece) > size and current:
            chunks.append("".join(current))
            tail, kept = [], 0
            for old in reversed(current):
                if kept + len(old) > overlap:
                    break
                tail.insert(0, old)
                kept += len(old)
            current, length = tail, kept
        current.append(piece)
        length += len(piece)
    if current:
        chunks.append("".join(current))
    return chunks or [""]


_INDEX_CACHE: "OrderedDict[str, list]" = OrderedDict()
_INDEX_LOCK = threading.Lock()


def _index_for(text: str) -> list:
    """[(chunk text, token count, term frequencies)] for one file's text. Chunking + tokenising a big file on every
    question is the slow part of retrieval on a weak CPU; the result only depends on the text, so it is cached."""
    key = hashlib.sha1(text.encode("utf-8", "ignore")).hexdigest()
    with _INDEX_LOCK:
        hit = _INDEX_CACHE.get(key)
        if hit is not None:
            _INDEX_CACHE.move_to_end(key)
            return hit
    built = []
    for piece in chunk_text(text):
        tokens = tokenize(piece)
        built.append((piece, len(tokens), Counter(tokens)))
    with _INDEX_LOCK:
        _INDEX_CACHE[key] = built
        while len(_INDEX_CACHE) > 48:
            _INDEX_CACHE.popitem(last=False)
    return built


def select_context(attachments: Sequence[Attachment], query: str, budget: int = 6000) -> tuple[str, bool]:
    """Text block for the prompt. Everything if it fits; otherwise the BM25-best chunks in file order.
    Returns (context, was_truncated)."""
    docs = [a for a in attachments if a.text.strip()]
    if not docs:
        return "", False
    if sum(len(a.text) for a in docs) <= budget:
        return "\n\n".join(f"===== FILE: {a.name} =====\n{a.text}\n===== END {a.name} =====" for a in docs), False

    chunks = []                                            # (file_index, part, parts, text, length, term_freq)
    for fi, doc in enumerate(docs):
        built = _index_for(doc.text)                       # tokenised once per file, then served from the cache
        for part, (piece, length, tf) in enumerate(built):
            chunks.append((fi, part, len(built), piece, length, tf))
    n = len(chunks)
    df: Counter = Counter()
    for chunk in chunks:
        df.update(chunk[5].keys())
    avg_len = sum(c[4] for c in chunks) / n or 1.0
    query_terms = set(tokenize(query))

    def bm25(chunk) -> float:
        length, tf = chunk[4], chunk[5]
        total = 0.0
        for term in query_terms:
            if term in tf:
                idf = math.log(1 + (n - df[term] + 0.5) / (df[term] + 0.5))
                total += idf * tf[term] * 2.5 / (tf[term] + 1.5 * (0.25 + 0.75 * length / avg_len))
        return total + 0.01 / (1 + chunk[1])               # tiny prior: beginnings win ties

    ranked = sorted(range(n), key=lambda i: bm25(chunks[i]), reverse=True)
    chosen, used = [], 0
    for i in ranked:
        cost = len(chunks[i][3]) + 60
        if used + cost > budget and chosen:
            continue
        chosen.append(i)
        used += cost
    chosen.sort(key=lambda i: (chunks[i][0], chunks[i][1]))

    out, last = [], None
    for i in chosen:
        fi, part, parts, text, *_ = chunks[i]
        if last is not None and not (last[0] == fi and last[1] + 1 == part):
            out.append("[...]")
        out.append(f"===== FILE: {docs[fi].name} (part {part + 1}/{parts}) =====\n{text}")
        last = (fi, part)
    return "\n\n".join(out), True


# ============================================================================
# 3. Tests: extraction, hold-out split, hard-coding detector
# ============================================================================
def _stdlib(module: str) -> bool:
    names = getattr(sys, "stdlib_module_names", ())
    return module.split(".")[0] in names


@functools.lru_cache(maxsize=256)
def _third_party(module: str) -> bool:
    """True when the module is a real installed library (lives in site-/dist-packages). Anything else that is not
    stdlib - `solution`, `main`, a local file - is treated as the code under test."""
    try:
        spec = importlib.util.find_spec(module.split(".")[0])
    except (ImportError, ValueError, AttributeError):
        return False
    return "-packages" in ((spec.origin or "") if spec else "").replace("\\", "/")


def tests_from_source(source: str) -> tuple[list[str], int]:
    """Extract EUTV-style test statements from Python source: module-level asserts and argument-free `def test_*`
    functions (each becomes one self-contained test). Imports of the solution module are dropped; stdlib imports
    are kept only where the test uses them. Returns (tests, skipped) - skipped = pytest fixtures/classes."""
    try:
        tree = ast.parse(source)
    except (SyntaxError, ValueError):
        return [], 0

    prefix: list[tuple[str, str]] = []                      # (bound name, line)
    for node in tree.body:
        if isinstance(node, ast.Import):
            for alias in node.names:
                bound = alias.asname or alias.name.split(".")[0]
                if _stdlib(alias.name) or _third_party(alias.name):
                    prefix.append((bound, ast.unparse(ast.Import(names=[alias]))))
                else:    # `import solution` -> expose the solution's namespace under that name
                    prefix.append((bound, f"{bound} = __import__('types').SimpleNamespace(**globals())"))
        elif (isinstance(node, ast.ImportFrom) and node.level == 0 and node.module
              and (_stdlib(node.module) or _third_party(node.module))):
            for alias in node.names:
                single = ast.ImportFrom(module=node.module, names=[alias], level=0)
                prefix.append((alias.asname or alias.name, ast.unparse(single)))

    def needed(node: ast.AST) -> list[str]:
        used = {n.id for n in ast.walk(node) if isinstance(n, ast.Name)}
        return [line for bound, line in prefix if bound in used or bound == "*"]

    tests, skipped = [], 0
    for node in tree.body:
        if isinstance(node, ast.Assert):
            tests.append("\n".join(needed(node) + [ast.unparse(node)]))
        elif isinstance(node, ast.FunctionDef) and node.name.startswith("test"):
            args = node.args
            required = len(args.posonlyargs) + len(args.args) - len(args.defaults)
            body = ast.unparse(node)
            if required > 0 or "pytest" in body:
                skipped += 1
                continue
            tests.append("\n".join(needed(node) + [body, f"{node.name}()"]))
        elif isinstance(node, ast.ClassDef) and node.name.startswith("Test"):
            skipped += 1
    return tests, skipped


def is_test_file(name: str, source: str, tests: Sequence[str]) -> bool:
    """A file is a test file if it is named like one, or contains tests and no non-test definitions."""
    base = os.path.basename(name).lower()
    if base.startswith("test_") or base.endswith("_test.py") or base == "tests.py":
        return True
    if not tests:
        return False
    try:
        tree = ast.parse(source)
    except (SyntaxError, ValueError):
        return False
    return not any(isinstance(n, (ast.FunctionDef, ast.ClassDef)) and not n.name.lower().startswith("test")
                   for n in tree.body)


def split_tests(tests: Sequence[str], holdout: float = 0.25) -> tuple[list[str], list[str]]:
    """Deterministic (hash-based) visible/hidden split. With fewer than 4 tests nothing is hidden: too little signal."""
    tests = list(tests)
    if len(tests) < 4 or holdout <= 0:
        return tests, []
    n_hidden = max(1, min(len(tests) - 2, round(len(tests) * holdout)))
    order = sorted(range(len(tests)), key=lambda i: hashlib.sha256(tests[i].encode()).hexdigest())
    hidden_idx = set(order[:n_hidden])
    return ([t for i, t in enumerate(tests) if i not in hidden_idx],
            [t for i, t in enumerate(tests) if i in hidden_idx])


def _distinctive_constants(tree: ast.AST) -> set:
    found = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Constant):
            v = node.value
            if (type(v) is int and abs(v) >= 10) or type(v) is float or (isinstance(v, (str, bytes)) and len(v) >= 3):
                found.add(v)
    return found


def hardcode_ratio(code: str, tests: Sequence[str]) -> float:
    """HEURISTIC: fraction of the tests' distinctive literals (numbers >= 10, floats, strings of 3+ chars) that also
    appear as constants in the solution. A solution that merely embeds expected values scores high; honest code that
    shares a few constants scores low. Returns 0.0 when there are too few literals to judge."""
    try:
        solution = _distinctive_constants(ast.parse(code))
    except (SyntaxError, ValueError):
        return 0.0
    literals: set = set()
    for test in tests:
        try:
            literals |= _distinctive_constants(ast.parse(test))
        except (SyntaxError, ValueError):
            continue
    if len(literals) < 3:
        return 0.0
    return len(literals & solution) / len(literals)


# ---- 3b. Test quality: lint, sandbox pre-flight, mutation check -------------------------------------------------------
def lint_tests(tests: Sequence[str]) -> list[str]:
    """Cheap static checks on the tests themselves (no model, no execution). A loop that optimises against broken
    tests only learns to satisfy the brokenness, so these are reported up front."""
    expected: dict[str, set] = {}
    seen: set[str] = set()
    duplicates = trivial = 0
    for test in tests:
        duplicates += test in seen
        seen.add(test)
        try:
            tree = ast.parse(test)
        except (SyntaxError, ValueError):
            continue
        for node in ast.walk(tree):
            if not isinstance(node, ast.Assert):
                continue
            body = node.test
            if not any(isinstance(n, (ast.Call, ast.Attribute, ast.Subscript, ast.Name)) for n in ast.walk(body)):
                trivial += 1
            if (isinstance(body, ast.Compare) and len(body.ops) == 1 and isinstance(body.ops[0], ast.Eq)
                    and isinstance(body.left, ast.Call)):
                try:
                    value = ast.literal_eval(body.comparators[0])
                except (ValueError, SyntaxError, TypeError):
                    continue
                expected.setdefault(ast.dump(body.left), set()).add(repr(value))
    out = []
    clashes = sum(1 for values in expected.values() if len(values) > 1)
    if clashes:
        out.append(f"{clashes} call(s) appear in the tests with different expected results: the tests contradict each "
                   "other, so no code can pass all of them")
    if trivial:
        out.append(f"{trivial} trivial assert(s) never exercise the code (e.g. `assert True`)")
    if duplicates:
        out.append(f"{duplicates} duplicate test(s): they inflate the pass rate without adding information")
    return out


_NETWORK_ROOTS = {"requests", "httpx", "urllib3", "aiohttp", "socket", "smtplib", "ftplib", "websockets", "openai",
                  "anthropic", "boto3", "grpc", "paramiko", "telnetlib"}
_NETWORK_FULL = {"urllib.request", "http.client", "xmlrpc.client", "http.server", "socketserver"}
_HEAVY_ROOTS = {"torch", "tensorflow", "transformers", "jax", "cv2", "sklearn", "scipy", "pandas", "matplotlib"}
_NETWORK_TEXT = re.compile(r"https?://|\brequests\.(?:get|post)|\bapi[ _-]?key\b|\bREST API\b|\bsocket\b", re.I)


def _imports(source: str) -> set[str]:
    found: set[str] = set()
    try:
        tree = ast.parse(source or "")
    except (SyntaxError, ValueError):
        return found
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            found.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.level == 0 and node.module:
            found.add(node.module)
            found.update(f"{node.module}.{alias.name}" for alias in node.names)
    return found


def _installed(root: str) -> bool:
    try:
        return importlib.util.find_spec(root) is not None
    except (ImportError, ValueError, AttributeError):
        return False


@dataclass(frozen=True)
class Preflight:
    warnings: tuple = ()
    addendum: str = ""          # appended to the task so the model writes code the sandbox can actually verify


def sandbox_preflight(task: str, tests: Sequence[str], initial_code: str = "") -> Preflight:
    """The sandbox has no network and a memory cap. Find out BEFORE spending iterations whether the code or the tests
    run into that, tell the user, and steer the model toward code that can be verified offline."""
    imported = _imports(initial_code)
    for test in tests:
        imported |= _imports(test)
    roots = {name.split(".")[0] for name in imported}
    warnings, notes = [], []

    network = sorted({r for r in roots if r in _NETWORK_ROOTS} | {n for n in imported if n in _NETWORK_FULL})
    if network or _NETWORK_TEXT.search(task or ""):
        if network:
            warnings.append(f"network modules ({', '.join(network)}) cannot connect inside the sandbox: only offline/"
                            "mocked behaviour can be verified")
        notes.append("The verification sandbox has NO network access. Never make real network calls at import time; "
                     "put any HTTP/socket access behind a small function or an injectable parameter so it can be "
                     "mocked, and keep the core logic testable offline.")

    heavy = sorted(r for r in roots if r in _HEAVY_ROOTS)
    missing_heavy = [r for r in heavy if not _installed(r)]
    if missing_heavy:
        warnings.append(f"{', '.join(missing_heavy)} is not installed in the sandbox environment: code that needs it "
                        "cannot be verified here")
    if heavy:
        notes.append("The sandbox has a small memory cap: prefer the standard library over large frameworks.")

    unknown = sorted(r for r in {n.split(".")[0] for n in _imports(initial_code)}
                     if not _stdlib(r) and not _installed(r) and r not in _HEAVY_ROOTS | _NETWORK_ROOTS)
    if unknown:
        warnings.append(f"your code imports {', '.join(unknown)}, which is not available in the sandbox: tests that "
                        "reach it will fail until it is inlined or replaced")
    return Preflight(tuple(warnings), ("\n\n" + " ".join(notes)) if notes else "")


_CMP_SWAP = {ast.Lt: ast.LtE, ast.LtE: ast.Lt, ast.Gt: ast.GtE, ast.GtE: ast.Gt, ast.Eq: ast.NotEq, ast.NotEq: ast.Eq}
_BIN_SWAP = {ast.Add: ast.Sub, ast.Sub: ast.Add, ast.Mult: ast.Add, ast.Div: ast.Mult, ast.FloorDiv: ast.Mult,
             ast.Mod: ast.Mult}
_BOOL_SWAP = {ast.And: ast.Or, ast.Or: ast.And}


class _Mutator(ast.NodeTransformer):
    """Applies exactly one small mutation: the `target`-th mutable site in traversal order. target=-1 only counts."""

    def __init__(self, target: int):
        self.target, self.sites = target, 0

    def _hit(self) -> bool:
        hit = self.sites == self.target
        self.sites += 1
        return hit

    def visit_Compare(self, node):
        self.generic_visit(node)
        node.ops = [_CMP_SWAP[type(op)]() if type(op) in _CMP_SWAP and self._hit() else op for op in node.ops]
        return node

    def visit_BinOp(self, node):
        self.generic_visit(node)
        if type(node.op) in _BIN_SWAP and self._hit():
            node.op = _BIN_SWAP[type(node.op)]()
        return node

    def visit_BoolOp(self, node):
        self.generic_visit(node)
        if type(node.op) in _BOOL_SWAP and self._hit():
            node.op = _BOOL_SWAP[type(node.op)]()
        return node

    def visit_Constant(self, node):
        if type(node.value) is int and self._hit():
            return ast.copy_location(ast.Constant(node.value + 1), node)
        return node


def make_mutants(code: str, limit: int) -> list[str]:
    """Up to `limit` distinct single-change variants of `code`, spread evenly over all mutable sites (deterministic)."""
    if limit <= 0 or len(code) > 20_000:
        return []
    try:
        tree = ast.parse(code)
    except (SyntaxError, ValueError):
        return []
    counter = _Mutator(-1)
    counter.visit(copy.deepcopy(tree))
    total = counter.sites
    if total == 0:
        return []
    indexes = range(total) if total <= limit else sorted({round(i * (total - 1) / max(1, limit - 1)) for i in range(limit)})
    base = ast.unparse(tree)
    out: list[str] = []
    for index in indexes:
        mutated = copy.deepcopy(tree)
        _Mutator(index).visit(mutated)
        ast.fix_missing_locations(mutated)
        source = ast.unparse(mutated)
        if source != base and source not in out:
            out.append(source)
    return out


def mutation_check(verifier, code: str, tests: Sequence[str], limit: int) -> tuple[int, int]:
    """(mutants that still pass EVERY test, mutants tried). Survivors mean the tests do not pin the behaviour down
    (a few may be equivalent mutants, so only a high ratio is treated as a signal)."""
    mutants = make_mutants(code, limit)
    if not mutants or not tests:
        return 0, 0
    results = verifier.run_many(mutants, [list(tests)] * len(mutants))
    return sum(1 for r in results if r.pass_rate >= 1.0), len(mutants)


# ============================================================================
# 4. SmartMemory: stronger retrieval, same file format as ExperienceMemory
# ============================================================================
def _canon(signature: str) -> str:
    text = re.sub(r"(['\"]).*?\1", "S", signature.lower())
    text = re.sub(r"0x[0-9a-f]+|\d+", "0", text)
    return re.sub(r"\s+", " ", text).strip()


def _ngrams(text: str, n: int = 3) -> Counter:
    padded = f" {text} "
    return Counter(padded[i:i + n] for i in range(max(1, len(padded) - n + 1)))


def _cosine(a: Counter, b: Counter) -> float:
    dot = sum(v * b.get(k, 0) for k, v in a.items())
    norm = math.sqrt(sum(v * v for v in a.values())) * math.sqrt(sum(v * v for v in b.values()))
    return dot / norm if norm else 0.0


def _exception_class(signature: str) -> str:
    match = re.search(r"([A-Za-z_]*(?:Error|Exception|Timeout))", signature)
    return match.group(1) if match else ""


class _HubSync:
    """OPTIONAL (off unless EIM_MEMORY_REPO is set): keeps the memory file in a private Hugging Face dataset repo.
    A Space's disk is wiped on every restart/sleep, so without this the engine forgets what it learned each time.
    Needs a write token in HF_TOKEN (Space secret). Failures never reach the user: memory stays a best-effort cache."""
    DELAY = 30.0                                  # seconds: pushes are batched, not sent after every lesson

    def __init__(self, repo: str, path: str, token: str | None):
        self.repo, self.path, self.token = repo, path, token
        self._timer: threading.Timer | None = None
        self._lock = threading.Lock()

    @classmethod
    def from_env(cls, path: str):
        repo = os.environ.get("EIM_MEMORY_REPO", "").strip()
        if not repo or not path:
            return None
        return cls(repo, path, os.environ.get("HF_TOKEN") or os.environ.get("HUGGING_FACE_HUB_TOKEN"))

    def pull(self) -> None:
        try:
            from huggingface_hub import hf_hub_download
            folder = os.path.dirname(os.path.abspath(self.path))
            os.makedirs(folder, exist_ok=True)
            hf_hub_download(repo_id=self.repo, repo_type="dataset", filename=os.path.basename(self.path),
                            token=self.token, local_dir=folder)
        except Exception:
            pass                                  # first run (nothing stored yet), no network, bad token ...

    def push_soon(self) -> None:
        with self._lock:
            if self._timer is not None:
                return
            self._timer = threading.Timer(self.DELAY, self._push)
            self._timer.daemon = True
            self._timer.start()

    def _push(self) -> None:
        with self._lock:
            self._timer = None
        try:
            from huggingface_hub import HfApi
            api = HfApi(token=self.token)
            api.create_repo(self.repo, repo_type="dataset", private=True, exist_ok=True)
            api.upload_file(path_or_fileobj=self.path, path_in_repo=os.path.basename(self.path), repo_id=self.repo,
                            repo_type="dataset", commit_message="EIM memory update")
        except Exception:
            pass


class SmartMemory(ExperienceMemory):
    """Drop-in replacement for ExperienceMemory (same JSON file). `hint` now:
       1. only compares failures of the same class (syntax / resource / zero / partial);
       2. normalises digits, quoted strings and spacing, so 'index 7' and 'index 12' look the same;
       3. blends char-trigram cosine, difflib ratio and an exception-class bonus;
       4. lets every close lesson VOTE for its strategy (weighted by similarity, newer lessons count more)."""

    def __init__(self, *args, **kwargs):
        path = kwargs.get("path", args[0] if args else "")
        self._hub = _HubSync.from_env(path if isinstance(path, str) else "")
        if self._hub is not None:
            self._hub.pull()                      # BEFORE the base class reads the file
        super().__init__(*args, **kwargs)

    def add(self, signature: str, strategy: str) -> None:
        super().add(signature, strategy)
        if getattr(self, "_hub", None) is not None:
            self._hub.push_soon()

    THRESHOLD = 0.55
    SCAN_LIMIT = 400            # newest lessons only: keeps the lookup O(1)-ish on a weak CPU however long the file grows

    def _features(self, signature: str):
        """(canonical form, trigram counter, exception class) - computed once per stored lesson, not once per query."""
        cache = self.__dict__.setdefault("_feature_cache", {})
        hit = cache.get(signature)
        if hit is None:
            canon = _canon(signature)
            hit = (canon, _ngrams(canon), _exception_class(signature))
            if len(cache) > 2 * self.SCAN_LIMIT:
                cache.clear()
            cache[signature] = hit
        return hit

    def hint(self, signature: str) -> str:
        with self._lock:
            items = list(self._items)[-self.SCAN_LIMIT:]
        if not items:
            return ""
        kind = signature.split("|", 1)[0]
        query, query_grams, query_exc = self._features(signature)
        votes: Counter = Counter()
        for rank, (stored, strategy) in enumerate(items):
            if stored.split("|", 1)[0] != kind:
                continue
            canon, grams, exc = self._features(stored)
            sim = 0.5 * _cosine(query_grams, grams) + 0.5 * difflib.SequenceMatcher(None, query, canon).ratio()
            if query_exc and query_exc == exc:
                sim = min(1.0, sim + 0.15)
            if sim >= self.THRESHOLD:
                votes[strategy] += sim * (1.0 + 0.5 * rank / len(items))
        return votes.most_common(1)[0][0] if votes else ""


# ============================================================================
# 5. EIMPlus: the loop with hold-out, restarts and "start from my code"
# ============================================================================
@dataclass(frozen=True)
class Report:
    best: Step
    first: Step
    visible_total: int
    hidden: ExecResult | None
    hidden_total: int
    hardcode: float
    restarts: int
    iterations: int
    warnings: tuple = ()


def restart_prompt(task: str, tests: Sequence[str], stuck_code: str, variant: int,
                   failure_feedback: str = "") -> str:
    angles = (
        "Use a completely different algorithm from the previous attempt.",
        "List the edge cases first (empty input, negatives, duplicates, large values), then write the simplest correct code.",
        "Ignore the previous approach entirely and re-derive the solution from the task statement.",
    )
    return (draft_prompt(task, tests)
            + f"\n\nA previous attempt got stuck:\n```python\n{stuck_code[:1500]}\n```\n{angles[variant % len(angles)]}"
            + (f"\n\nDiagnostic feedback from the failed tests:\n{failure_feedback[:1800]}"
               if failure_feedback else "")
            + "\nDo not repeat the same failed approach or merely rename variables. Use the diagnostic feedback to identify "
              "the underlying bug, then choose a meaningfully different implementation if the previous approach stalled. "
              "Write a general solution for every valid input: never embed expected test values, special-case test inputs, "
              "or build lookup tables of answers. Preserve the required function names and signatures; handle empty input, "
              "negative values, duplicates and boundaries; prefer the standard library; make no network calls at import time. "
              "Return the complete corrected code, not an explanation.")


_STOP = frozenset("""assert true false none def return import from the and for that with this are you your was were
    will can could should would must not but all any each per into than then them they have has had its our out one two
    write make create implement function code python test tests please given input output value values result results
    using use used need needs want wants get gets set sets list item items""".split())


def _content_tokens(text: str) -> list[str]:
    """Words that identify a PROBLEM: no numbers, no 1-2 letter names, no boilerplate such as `assert` or `return`.
    Without this, any two tasks sharing 'assert', 'a', 'b' and a few digits look 'similar'."""
    return [t for t in tokenize(text) if len(t) >= 3 and not t.isdigit() and t not in _STOP]


class RepairLog:
    """VERIFIED experience: what was broken, what fixed it, which tests proved it. One JSON line per record.

    Only runs that passed every visible AND hidden test are stored, and runs that look like hard-coding are never stored.
    Records whose tests looked weak (mutation check) or contradictory are kept but flagged `weak`: they are neither
    retrieved nor exported for training. `similar()` finds the closest earlier problems so the model sees a worked,
    verified example - not just a generic "try strategy X". Off when EIM_REPAIR_LOG=0 or when there is no memory path."""
    MAX_RECORDS = 500
    MIN_SIM = 0.30
    SCAN = 300                       # newest records searched per query

    def __init__(self, path: str):
        self.path = path
        self._lock = threading.Lock()
        self._hub = _HubSync.from_env(path)
        if self._hub is not None:
            self._hub.pull()
        self._records: list[dict] = self._load()
        self._grams: dict[str, Counter] = {}

    @classmethod
    def for_config(cls, cfg) -> "RepairLog | None":
        if os.environ.get("EIM_REPAIR_LOG", "1") == "0":
            return None
        explicit = os.environ.get("EIM_REPAIR_LOG_PATH", "").strip()
        base = getattr(cfg, "memory_path", "") or ""
        if explicit:
            return cls(explicit)
        if base:
            return cls(os.path.join(os.path.dirname(os.path.abspath(base)), "eim_repairs.jsonl"))
        return None

    def _load(self) -> list[dict]:
        out: list[dict] = []
        try:
            with open(self.path, encoding="utf-8") as handle:
                for line in handle:
                    try:
                        record = json.loads(line)
                    except ValueError:
                        continue
                    if isinstance(record, dict) and record.get("id") and record.get("after"):
                        out.append(record)
        except OSError:
            pass
        return out[-self.MAX_RECORDS:]

    def __len__(self) -> int:
        return len(self._records)

    def add(self, *, task: str, tests: Sequence[str], signature: str, before: str, after: str, strategy: str,
            weak: bool = False) -> bool:
        key = hashlib.sha1((task + "\0" + after).encode("utf-8", "ignore")).hexdigest()[:16]
        repaired = bool(before) and before != after
        record = {"id": key, "kind": "repair" if repaired else "solution", "task": task[:1500],
                  "tests": [t[:300] for t in list(tests)[:8]], "signature": signature[:300] if repaired else "",
                  "before": before[:4000] if repaired else "", "after": after[:4000], "strategy": strategy,
                  "weak": bool(weak), "ts": int(time.time())}
        with self._lock:
            if any(r["id"] == key for r in self._records):
                return False
            self._records.append(record)
            try:
                folder = os.path.dirname(os.path.abspath(self.path))
                os.makedirs(folder, exist_ok=True)
                if len(self._records) > self.MAX_RECORDS + 100:          # compact: keep the newest MAX_RECORDS
                    self._records = self._records[-self.MAX_RECORDS:]
                    with open(self.path, "w", encoding="utf-8") as handle:
                        handle.writelines(json.dumps(r, ensure_ascii=False) + "\n" for r in self._records)
                else:
                    with open(self.path, "a", encoding="utf-8") as handle:
                        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
            except OSError:
                pass                                                        # memory is best-effort
        if self._hub is not None:
            self._hub.push_soon()
        return True

    def _features(self, record: dict) -> Counter:
        hit = self._grams.get(record["id"])
        if hit is None:
            hit = Counter(_content_tokens(record["task"] + " " + " ".join(record["tests"])))
            if len(self._grams) > 2 * self.MAX_RECORDS:
                self._grams.clear()
            self._grams[record["id"]] = hit
        return hit

    def similar(self, task: str, tests: Sequence[str], signature: str | None = None, k: int = 2) -> list[dict]:
        with self._lock:
            records = list(self._records)[-self.SCAN:]
        if not records:
            return []
        query = Counter(_content_tokens(task + " " + " ".join(tests)))
        want_class = _exception_class(signature) if signature else ""
        scored = []
        for rank, record in enumerate(records):
            if record.get("weak"):
                continue
            sim = _cosine(query, self._features(record))
            if want_class and record["kind"] == "repair" and _exception_class(record["signature"]) == want_class:
                sim = min(1.0, sim + 0.1)
            if sim >= self.MIN_SIM:
                scored.append((sim + 0.001 * rank, record))              # newer wins ties
        scored.sort(key=lambda t: t[0], reverse=True)
        return [r for _, r in scored[:k]]

    @staticmethod
    def render(records: Sequence[dict], max_chars: int = 1600) -> str:
        if not records:
            return ""
        out = ("\n\nVerified reference examples from earlier runs on similar problems (adapt the idea, "
               "do not copy blindly):\n")
        for number, record in enumerate(records, 1):
            parts = [f"### Example {number}\nTask: {record['task'][:300]}\n"]
            if record["kind"] == "repair":
                parts.append(f"Broken version:\n```python\n{record['before'][:700]}\n```\n")
            parts.append(f"Verified solution:\n```python\n{record['after'][:900]}\n```\n")
            piece = "".join(parts)
            if len(out) + len(piece) > max_chars and number > 1:
                break
            out += piece
        return out[:max_chars + 400]


def export_sft(log: RepairLog, out_path: str) -> int:
    """Verified, non-weak records -> chat-format JSONL (one {"messages": [...]} per line) for LoRA/QLoRA SFT, e.g. with
    TRL's SFTTrainer. Returns the number of examples written. Hold out a fixed evaluation set BEFORE training on these
    (eval_eim.py) and compare the model before/after; do not retrain after every single run."""
    written = 0
    with open(out_path, "w", encoding="utf-8") as handle:
        for record in log._records:
            if record.get("weak"):
                continue
            tests = "\n".join(record["tests"])
            if record["kind"] == "repair":
                prompt = (f"Fix the code so that every test passes.\n\nTask:\n{record['task']}\n\nTests:\n{tests}\n\n"
                          f"Code:\n```python\n{record['before']}\n```")
            else:
                prompt = f"Write Python code for the task. It must pass the tests.\n\nTask:\n{record['task']}\n\nTests:\n{tests}"
            answer = f"```python\n{record['after']}\n```"
            handle.write(json.dumps({"messages": [{"role": "user", "content": prompt},
                                                   {"role": "assistant", "content": answer}]}, ensure_ascii=False) + "\n")
            written += 1
    return written


_RESULT_CACHE: "OrderedDict[str, tuple]" = OrderedDict()
_RESULT_LOCK = threading.Lock()
_RESULT_CACHE_MAX = 16


def _cache_enabled() -> bool:
    return os.environ.get("EIM_RESULT_CACHE", "1") != "0"


class EIMPlus:
    STALE_BEFORE_RESTART = 2
    HARDCODE_LIMIT = 0.5
    MUTATION_SURVIVAL_LIMIT = 0.5

    def __init__(self, lm: LanguageModel, verifier: EUTV | None = None, cfg: Config = CFG,
                 memory: ExperienceMemory | None = None, profile: Profile | None = None,
                 repair_log: "RepairLog | None | str" = "auto"):
        self.cfg = cfg
        self.lm = lm
        self.verifier = verifier or EUTV(cfg)
        self.ifc = IFC(cfg)
        self.memory = memory if memory is not None else SmartMemory(path=cfg.memory_path)
        self.dcme = DCME(lm, self.memory, cfg)
        self.profile = profile or hardware_profile()
        # verified-experience log: "auto" = derived from cfg.memory_path (none when that is empty), None = off
        self.repairs = RepairLog.for_config(cfg) if repair_log == "auto" else repair_log

    def _examples(self, task: str, tests: Sequence[str], signature: str | None) -> str:
        if self.repairs is None:
            return ""
        try:
            lean = self.profile.name == "lean"
            found = self.repairs.similar(task, tests, signature, k=1 if lean else 2)
            return RepairLog.render(found, 1000 if lean else 1600)
        except Exception:
            return ""

    # ------------------------------------------------------------------ public
    def run(self, task: str, tests: Sequence[str], initial_code: str | None = None, max_iter: int = 4, k: int = 3,
            holdout: float = 0.25, *, mutation: bool | None = None,
            time_budget: float | None = None) -> Iterator[Step | Report]:
        """Same contract as before (yields Steps, then one Report). New, optional, keyword-only:
             mutation     - run the mutation check on the final code (default: on)
             time_budget  - seconds after which no NEW round starts and the best version is returned
                            (default: env EIM_TIME_BUDGET, 0 = unlimited)
        A run that fully succeeded is replayed instantly when the very same request comes again."""
        tests = list(tests)
        raw = json.dumps([task, tests, initial_code or "", int(max_iter), int(k), float(holdout), id(self.lm),
                          getattr(self.cfg, "model_name", ""), getattr(self.cfg, "max_new_tokens", None),
                          self.profile.name, _repair_rules()], ensure_ascii=False)
        key = hashlib.sha256(raw.encode("utf-8", "ignore")).hexdigest()
        if _cache_enabled():
            with _RESULT_LOCK:
                hit = _RESULT_CACHE.get(key)
                if hit is not None:
                    _RESULT_CACHE.move_to_end(key)
            if hit is not None:
                yield from hit
                return
        events: list = []
        for event in self._run(task, tests, initial_code, int(max_iter), int(k), float(holdout), mutation, time_budget):
            if isinstance(event, Report):
                if _cache_enabled() and event.best.result.pass_rate == 1.0 and (
                        event.hidden is None or event.hidden.pass_rate == 1.0):
                    with _RESULT_LOCK:                       # stored BEFORE the Report is yielded: callers may stop there
                        _RESULT_CACHE[key] = tuple(events) + (event,)
                        while len(_RESULT_CACHE) > _RESULT_CACHE_MAX:
                            _RESULT_CACHE.popitem(last=False)
                yield event
                return
            events.append(event)
            yield event

    # ------------------------------------------------------------------ the loop
    def _run(self, task: str, tests: list[str], initial_code: str | None, max_iter: int, k: int, holdout: float,
             mutation: bool | None, time_budget: float | None) -> Iterator[Step | Report]:
        profile = self.profile
        budget = effective_time_budget(profile, time_budget)
        deadline = time.monotonic() + budget if budget > 0 else None
        visible, hidden = split_tests(tests, holdout)

        # --- before spending any compute: is the problem itself sound, and can this sandbox verify it?
        warnings: list[str] = list(lint_tests(tests))
        pre = sandbox_preflight(task, tests, initial_code or "")
        warnings += pre.warnings
        original_task = task
        base_task = task + pre.addendum
        task = base_task + self._examples(original_task, tests, None)

        asked = k
        k = max(1, min(k, profile.k_cap))
        if k < asked:
            warnings.append(f"candidates per round reduced from {asked} to {k} to fit this machine "
                            f"({profile.name} profile; EIM_PROFILE=full lifts the cap)")
        # ladder: ONE candidate first (usually enough), the remaining ones only if it did not help
        batches = [1, k - 1] if profile.ladder and k >= 2 else [k]

        if initial_code and initial_code.strip():
            draft, label = normalise(initial_code), "given"
        else:
            raw = self.lm.generate([draft_prompt(task, visible)], 0.2, self.cfg.max_new_tokens)[0]
            draft, label = normalise(extract_code(raw)), "draft"
        result = self.verifier.run(draft, visible)
        first = best = Step(0, label, True, label, draft, result, score(draft, result, self.cfg))
        yield best
        if first.result.pass_rate < 1.0:             # now the failure is known: look for repairs of the same kind
            task = base_task + self._examples(original_task, tests, self.ifc.signature(first.result))

        tried, stale, restarts, ran, extended = {draft}, 0, 0, 0, 0
        iteration, limit, bonus = 0, max_iter, profile.bonus_iters
        while iteration < limit:
            iteration += 1
            # An unresolved failure must get its promised restart before the convergence guard
            # can stop the loop. Otherwise a verifier that treats repeated stale rounds as
            # convergence can make the restart branch unreachable.
            stuck = stale >= self.STALE_BEFORE_RESTART and best.result.pass_rate < 1.0
            if self.ifc.converged(best.result, stale) and not stuck:
                break
            if deadline is not None and time.monotonic() >= deadline:
                warnings.append(f"time budget of {budget:.0f}s reached: returning the best version found so far")
                break
            ran = iteration
            # Keep diagnostics even during a restart. Previously the stuck branch erased them,
            # forcing the model to guess why its earlier code failed.
            feedback = self.ifc.feedback(best.code, best.result, best.score)
            if stuck:
                restarts += 1
            temperature = min(0.9, 0.3 + 0.2 * stale)

            accepted, fallback, offset = False, None, 0
            for batch_k in batches:
                if stuck:
                    prompts = [restart_prompt(task, visible, best.code, offset + i, feedback) for i in range(batch_k)]
                    outputs = self.lm.generate(prompts, 0.9, self.cfg.max_new_tokens)
                    mutations = []
                    for text in outputs:
                        candidate = normalise(extract_code(text))
                        if candidate and candidate not in tried:
                            tried.add(candidate)
                            mutations.append(Mutation("restart", candidate))
                else:
                    mutations = self.dcme.propose(task, best.code, feedback, best.result, batch_k,
                                                  min(0.9, temperature + (0.1 if offset else 0.0)), tried)
                offset += batch_k
                if not mutations:
                    continue

                results = self.verifier.run_many([m.code for m in mutations], [visible] * len(mutations))
                scored = [(score(m.code, r, self.cfg), m, r) for m, r in zip(mutations, results)]
                eligible = [t for t in scored if t[2].pass_rate >= best.result.pass_rate]   # never regress
                sc, mutation_, res = max(eligible or scored, key=lambda t: t[0].reward)

                if eligible and sc.reward > best.score.reward + self.cfg.min_gain:
                    if res.pass_rate > best.result.pass_rate and mutation_.strategy in STRATEGIES:
                        self.memory.add(self.ifc.signature(best.result), mutation_.strategy)
                    best = Step(iteration, "accepted", True, mutation_.strategy, mutation_.code, res, sc)
                    accepted = True
                    break
                if fallback is None or sc.reward > fallback[0].reward:
                    fallback = (sc, mutation_, res)

            if accepted:
                stale = 0
                yield best
            elif fallback is not None:
                stale += 1
                sc, mutation_, res = fallback
                yield Step(iteration, "rejected", False, mutation_.strategy, mutation_.code, res, sc)
            else:
                stale += 1
                yield Step(iteration, "no new candidate", False, "-", best.code, best.result, best.score)

            # adaptive budget: the last round still improved the code -> one more round, within the profile's bound
            if iteration == limit and bonus > 0 and accepted and best.result.pass_rate < 1.0:
                limit, bonus, extended = limit + 1, bonus - 1, extended + 1

        if extended:
            warnings.append(f"the iteration budget was extended by {extended} round(s) because the code was still "
                            "improving (adaptive budget)")

        # --- final checks on tests the loop never selected on
        hidden_result = self.verifier.run(best.code, hidden) if hidden else None
        ratio = hardcode_ratio(best.code, tests)
        if hidden_result is not None and hidden_result.pass_rate < 1.0 and best.result.pass_rate == 1.0:
            warnings.append(f"passes every visible test but only {hidden_result.passed}/{hidden_result.total} "
                            "hidden ones: likely overfitting, add more diverse tests")
        if ratio >= self.HARDCODE_LIMIT:
            warnings.append(f"{ratio:.0%} of the tests' distinctive values are embedded in the code: "
                            "it may be hard-coding answers")
        if not hidden:
            warnings.append("no hidden tests (needs at least 4 tests): generalisation was not measured")

        perfect = best.result.pass_rate == 1.0 and (hidden_result is None or hidden_result.pass_rate == 1.0)
        if (mutation is None or mutation) and perfect and len(tests) >= 2 \
                and (deadline is None or time.monotonic() < deadline):
            survived, tried_n = mutation_check(self.verifier, best.code, tests, profile.mutants)
            if tried_n >= 3 and survived / tried_n >= self.MUTATION_SURVIVAL_LIMIT:
                warnings.append(f"the tests look weak: {survived}/{tried_n} small deliberate bugs injected into the "
                                "final code still pass every test (a few may be harmless equivalents). "
                                "Add edge-case tests (empty input, negatives, boundaries)")
        if self.repairs is not None and perfect and ratio < self.HARDCODE_LIMIT and len(tests) >= 2:
            weak = any(("tests look weak" in w) or ("contradict" in w) for w in warnings)
            broken = first.result.pass_rate < 1.0
            try:
                self.repairs.add(task=original_task, tests=tests, signature=self.ifc.signature(first.result) if broken else "",
                                 before=first.code if broken else "", after=best.code, strategy=best.strategy, weak=weak)
            except Exception:
                pass                                                      # never let bookkeeping break an answer
        yield Report(best, first, len(visible), hidden_result, len(hidden), ratio, restarts, ran, tuple(warnings))


# ============================================================================
# 6. Self-test: python eim_plus.py --selftest
# ============================================================================
def run_selftest() -> int:
    import shutil
    import zipfile as zf
    from app import _ScriptedLM

    checks = 0

    def check(name: str, condition: bool) -> None:
        nonlocal checks
        checks += 1
        print(("PASS " if condition else "FAIL ") + name)
        if not condition:
            raise SystemExit(1)

    tmp = tempfile.mkdtemp(prefix="eimplus_")
    try:
        def write(name: str, data: bytes | str) -> str:
            path = os.path.join(tmp, name)
            with open(path, "wb") as handle:
                handle.write(data if isinstance(data, bytes) else data.encode("utf-8"))
            return path

        # ---- attachments
        check("utf-8 text", read_attachment(write("a.txt", "مرحبا بالعالم")).text == "مرحبا بالعالم")
        check("cp1256 (Windows Arabic) text", "مرحبا" in read_attachment(write("b.txt", "مرحبا".encode("cp1256"))).text)
        check("code file kind", read_attachment(write("c.py", "x = 1\n")).kind == "code")
        check("binary is rejected", read_attachment(write("d.txt", b"\x00\x01\x02")).kind == "unsupported")
        nb = json.dumps({"cells": [{"cell_type": "code", "source": ["print(1)\n"]}, {"cell_type": "markdown", "source": "hi"}]})
        check("ipynb cells", "print(1)" in read_attachment(write("e.ipynb", nb)).text)
        check("csv table", read_attachment(write("f.csv", "a,b\n1,2\n")).kind == "table")

        from docx import Document
        doc = Document()
        doc.add_paragraph("Quarterly revenue grew 12 percent")
        doc.save(os.path.join(tmp, "g.docx"))
        check("docx text (zip+xml parser)", "revenue grew" in read_attachment(os.path.join(tmp, "g.docx")).text)

        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["name", "score"])
        wb.active.append(["ali", 91])
        wb.save(os.path.join(tmp, "h.xlsx"))
        check("xlsx rows", "ali,91" in read_attachment(os.path.join(tmp, "h.xlsx")).text)

        with zf.ZipFile(os.path.join(tmp, "i.zip"), "w") as archive:
            archive.writestr("pkg/mod.py", "def hello():\n    return 1\n")
            archive.writestr("../evil.py", "print('x')")
            archive.writestr("img.bin", b"\x00\x01")
        zip_att = read_attachment(os.path.join(tmp, "i.zip"))
        check("zip: reads text members, never extracts", "def hello" in zip_att.text and not os.path.exists(os.path.join(tmp, "..", "evil.py")))

        # Generate a standards-compliant tiny PDF with correct xref offsets. A malformed
        # handcrafted PDF made pypdf spend unbounded time recovering from startxref=0.
        pdf_objects = [
            b"<</Type/Catalog/Pages 2 0 R>>",
            b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
            b"<</Length 49>>\nstream\nBT /F1 12 Tf 20 100 Td (Hello PDF world) Tj ET\nendstream",
            b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
        ]
        pdf_parts = [b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"]
        offsets = [0]
        for object_id, payload in enumerate(pdf_objects, 1):
            offsets.append(sum(map(len, pdf_parts)))
            pdf_parts.append(f"{object_id} 0 obj\n".encode() + payload + b"\nendobj\n")
        xref_offset = sum(map(len, pdf_parts))
        pdf_parts.append(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
        for offset in offsets[1:]:
            pdf_parts.append(f"{offset:010d} 00000 n \n".encode())
        pdf_parts.append(f"trailer\n<</Root 1 0 R/Size {len(offsets)}>>\nstartxref\n{xref_offset}\n%%EOF\n".encode())
        pdf = b"".join(pdf_parts)
        check("pdf text", "Hello PDF world" in read_attachment(write("j.pdf", pdf)).text)

        from PIL import Image
        Image.new("RGB", (32, 16), "red").save(os.path.join(tmp, "k.png"))
        image_att = read_attachment(os.path.join(tmp, "k.png"))
        check("image size", image_att.kind == "image" and image_att.image_size == (32, 16))
        check("fake image is rejected gracefully", read_attachment(write("l.png", "not an image")).kind == "unsupported")
        check("entity bomb is refused", "could not read" in read_attachment(
            _bomb_docx(tmp)).note)

        # ---- retrieval
        filler = "\n".join(f"line {i} about nothing in particular" for i in range(600))
        secret = "The refund policy allows returns within 30 days of purchase."
        big = Attachment("policy.txt", "", "text", text=filler + "\n" + secret + "\n" + filler)
        context, truncated = select_context([big], "what is the refund policy", budget=3000)
        check("BM25 finds the relevant chunk in a big file", truncated and secret in context and len(context) < 4500)
        arabic = Attachment("ar.txt", "", "text", text=filler + "\nسياسة الاسترجاع تسمح بإرجاع المنتج خلال ثلاثين يوما\n" + filler)
        context, _ = select_context([arabic], "ما هي سياسه الاسترجاع", budget=3000)
        check("Arabic retrieval (normalised)", "ثلاثين يوما" in context)
        small, truncated = select_context([Attachment("s.txt", "", "text", text="tiny")], "q")
        check("small files are included whole", not truncated and "tiny" in small)

        # ---- tests extraction
        src = ("import math\nfrom solution import area\n\ndef test_circle():\n    assert abs(area(1) - math.pi) < 1e-9\n\n"
               "def test_fixture(tmp_path):\n    pass\n\nassert area(0) == 0\n")
        extracted, skipped = tests_from_source(src)
        check("test extraction: asserts + test functions, fixtures skipped", len(extracted) == 2 and skipped == 1)
        check("plain asserts stay single statements (rich feedback)", extracted[1] == "assert area(0) == 0")
        check("stdlib imports are kept where used", "import math" in extracted[0])
        check("is_test_file by name/content", is_test_file("test_x.py", "", []) and not is_test_file("sol.py", "def f(): pass", []))
        aliased, _ = tests_from_source("import solution as s\nassert s.f(1) == 1\n")
        eutv_probe = EUTV(Config(memory_path="", per_test_seconds=1, wall_seconds=3.0, cpu_seconds=2))
        check("`import solution` alias works in the sandbox",
              eutv_probe.run("def f(x):\n    return x", aliased).pass_rate == 1.0)

        # ---- split + hardcode
        many = [f"assert f({i}) == {i * i}" for i in range(8)]
        vis, hid = split_tests(many, 0.25)
        check("hold-out split is deterministic and disjoint", len(hid) == 2 and not set(vis) & set(hid) and split_tests(many, 0.25) == (vis, hid))
        check("fewer than 4 tests: nothing hidden", split_tests(many[:3], 0.25)[1] == [])
        cheat = "def f(x):\n    return {10: 100, 20: 400, 30: 900}[x]"
        honest = "def f(x):\n    return x * x"
        big_tests = ["assert f(10) == 100", "assert f(20) == 400", "assert f(30) == 900"]
        check("hard-coding detector flags lookup tables", hardcode_ratio(cheat, big_tests) >= 0.5)
        check("hard-coding detector leaves honest code alone", hardcode_ratio(honest, big_tests) < 0.5)

        # ---- smart memory
        memory = SmartMemory(path="")
        memory.add("partial|IndexError: list index out of range", "repair")
        memory.add("zero|NameError: name 'x' is not defined", "rewrite")
        check("memory: same class, different details", memory.hint("partial|IndexError: list index out of range at 12") == "repair")
        check("memory: never crosses failure classes", memory.hint("zero|IndexError: list index out of range") == "")
        memory.add("partial|AssertionError: expected 3", "rewrite")
        memory.add("partial|AssertionError: expected 5", "rewrite")
        memory.add("partial|AssertionError: expected 7", "repair")
        check("memory: similar lessons vote", memory.hint("partial|AssertionError: expected 9") == "rewrite")

        # ---- the loop
        eutv = EUTV(Config(per_test_seconds=1, wall_seconds=3.0, cpu_seconds=2, memory_path=""))
        fast = Config(per_test_seconds=1, wall_seconds=3.0, cpu_seconds=2, memory_path="")

        # (a) start from the user's own buggy code
        buggy = "def add(a, b):\n    return a - b"
        fixed = "def add(a, b):\n    return a + b"
        events = list(EIMPlus(_ScriptedLM([fixed]), eutv, fast, SmartMemory(path="")).run(
            "add numbers", ["assert add(1, 2) == 3", "assert add(5, 5) == 10"], initial_code=buggy, max_iter=3, k=2))
        steps = [e for e in events if isinstance(e, Step)]
        check("loop starts from the given code", steps[0].label == "given" and steps[0].result.pass_rate < 1.0)
        check("loop repairs the given code", steps[-1].result.pass_rate == 1.0 and isinstance(events[-1], Report))

        # (b) overfitting is caught by hidden tests
        square_tests = [f"assert sq({i}) == {i * i}" for i in (2, 3, 4, 5, 6, 7, 8, 9)]
        v, h = split_tests(square_tests, 0.25)
        table = {int(re.search(r"sq\((\d+)\)", t).group(1)): int(t.rsplit("==", 1)[1]) for t in v}
        lookup = f"def sq(n):\n    return {table}[n]"
        report = list(EIMPlus(_ScriptedLM([lookup]), eutv, fast, SmartMemory(path="")).run(
            "square", square_tests, max_iter=2, k=1))[-1]
        check("overfitting: visible 100%, hidden < 100%, warning raised",
              report.best.result.pass_rate == 1.0 and report.hidden.pass_rate < 1.0
              and any("overfitting" in w for w in report.warnings))

        # (c) escape a local minimum with a restart
        right = "def add(a, b):\n    return a + b"
        script = ["def add(a, b):\n    return 0", "def add(a, b):\n    return 0 * a", "def add(a, b):\n    return 0 * b", right]
        events = list(EIMPlus(_ScriptedLM(script), eutv, fast, SmartMemory(path="")).run(
            "add", ["assert add(1, 2) == 3", "assert add(2, 2) == 4"], max_iter=5, k=1))
        report = events[-1]
        check("stuck loop restarts and escapes", report.restarts >= 1 and report.best.result.pass_rate == 1.0)

        # ---- hardware profile, cache, third-party imports in tests
        check("hardware profile is detected", hardware_profile().name in _PROFILES and hardware_profile().k_cap >= 1)
        before = select_context([big], "refund policy", budget=3000)
        check("retrieval index is cached and gives identical results",
              len(_INDEX_CACHE) > 0 and select_context([big], "refund policy", budget=3000) == before)
        third, _ = tests_from_source("import pypdf\nassert pypdf.__name__ == 'pypdf'\n")
        check("installed third-party imports are kept in extracted tests", "import pypdf" in third[0])

        # ---- test quality (lint) and sandbox pre-flight
        found = lint_tests(["assert f(1) == 2", "assert f(1) == 3", "assert True", "assert f(1) == 2"])
        check("lint: contradictory, trivial and duplicate tests are reported",
              any("contradict" in w for w in found) and any("trivial" in w for w in found)
              and any("duplicate" in w for w in found))
        check("lint: clean tests stay silent", lint_tests(["assert f(1) == 2", "assert f(2) == 4"]) == [])
        pf = sandbox_preflight("get data from https://api.example.com", ["assert f() == 1"],
                               "import requests\ndef f():\n    return requests.get('x')")
        check("pre-flight: network code is flagged and the model is told", bool(pf.warnings) and "NO network" in pf.addendum)
        pf = sandbox_preflight("add", ["assert add(1, 2) == 3"], "import definitely_not_installed_xyz\n")
        check("pre-flight: an unavailable import is flagged", any("definitely_not_installed_xyz" in w for w in pf.warnings))
        pf = sandbox_preflight("add", ["assert add(1, 2) == 3"], "import json\n")
        check("pre-flight: clean problems stay silent", pf.warnings == () and pf.addendum == "")

        # ---- mutation check
        clamp = "def clamp(x):\n    if x > 5:\n        return 5\n    return x"
        weak = ["assert clamp(0) == 0", "assert clamp(1) == 1", "assert clamp(2) == 2"]
        strong = weak + ["assert clamp(5) == 5", "assert clamp(6) == 5", "assert clamp(9) == 5", "assert clamp(-3) == -3"]
        check("mutants are deterministic and differ from the original",
              make_mutants(clamp, 5) == make_mutants(clamp, 5) and clamp not in make_mutants(clamp, 5)
              and len(make_mutants(clamp, 5)) == 3)
        lean = Profile("lean", 2, 4.0, False, 3, True, 3, 1)
        flat = Profile("test", 4, 8.0, False, 8, False, 0, 0)        # no ladder, no mutation check, no bonus rounds
        rep = list(EIMPlus(_ScriptedLM([clamp]), eutv, fast, SmartMemory(path=""), profile=lean).run(
            "clamp to 5", weak, max_iter=1, k=1))[-1]
        check("mutation check: weak tests are reported", any("tests look weak" in w for w in rep.warnings))
        rep = list(EIMPlus(_ScriptedLM([clamp]), eutv, fast, SmartMemory(path=""), profile=lean).run(
            "clamp to 5", strong, max_iter=1, k=1))[-1]
        check("mutation check: strong tests are not reported", not any("tests look weak" in w for w in rep.warnings))

        # ---- RepairLog: verified experience, retrieval, export
        rl_path = os.path.join(tmp, "eim_repairs.jsonl")
        rlog = RepairLog(rl_path)
        add_t = ["assert add(1, 2) == 3", "assert add(2, 2) == 4", "assert add(0, 5) == 5", "assert add(-1, 1) == 0"]
        bad_add = "def add(a, b):\n    return a - b"
        eng = EIMPlus(_ScriptedLM([right]), eutv, fast, SmartMemory(path=""), profile=flat, repair_log=rlog)
        rep = list(eng.run("write add(a, b) that sums two numbers", add_t, bad_add, 3, 1))[-1]
        check("RepairLog: a verified repair is stored with before/after", len(rlog) == 1
              and rlog._records[0]["kind"] == "repair" and rlog._records[0]["before"] == normalise(bad_add)
              and rlog._records[0]["after"] == rep.best.code)
        check("RepairLog: survives a restart (read back from disk)", len(RepairLog(rl_path)) == 1)
        check("RepairLog: the same solution is not stored twice", not rlog.add(
            task="write add(a, b) that sums two numbers", tests=add_t, signature="", before=bad_add, after=rep.best.code,
            strategy="x"))
        seen_prompts: list[str] = []

        class _SpyLM(_ScriptedLM):
            def generate(self, prompts, temperature, max_new_tokens):
                seen_prompts.extend(prompts)
                return super().generate(prompts, temperature, max_new_tokens)
        eng2 = EIMPlus(_SpyLM([right]), eutv, fast, SmartMemory(path=""), profile=flat, repair_log=rlog)
        list(eng2.run("write add(a, b) that sums two numbers please", add_t, None, 2, 1))
        check("RepairLog: a similar earlier problem is shown to the model as a verified example",
              any("Verified reference examples" in prompt and "return a + b" in prompt for prompt in seen_prompts[:1]))
        unrelated = rlog.similar("parse a csv file and compute median of columns", ["assert median([1, 2, 3]) == 2"])
        check("RepairLog: unrelated problems get no example", unrelated == [])
        weak_log = RepairLog(os.path.join(tmp, "weak.jsonl"))
        weak_log.add(task="clamp value to 5", tests=weak, signature="", before="", after=clamp, strategy="draft", weak=True)
        check("RepairLog: weak-test records are never retrieved or exported",
              weak_log.similar("clamp value to 5", weak) == []
              and export_sft(weak_log, os.path.join(tmp, "weak_sft.jsonl")) == 0)
        out_file = os.path.join(tmp, "sft.jsonl")
        exported = export_sft(rlog, out_file)
        lines = [json.loads(line) for line in open(out_file, encoding="utf-8")]
        check("export_sft: chat-format JSONL with the verified fix as the answer",
              exported == len(rlog) == 2 and all("return a + b" in row["messages"][1]["content"] for row in lines)
              and any("Fix the code" in row["messages"][0]["content"] for row in lines))
        hard = EIMPlus(_ScriptedLM(["def f(x):\n    return {10: 100, 20: 400, 30: 900, 40: 1600}[x]"]), eutv, fast,
                       SmartMemory(path=""), profile=flat, repair_log=RepairLog(os.path.join(tmp, "hard.jsonl")))
        list(hard.run("square", ["assert f(10) == 100", "assert f(20) == 400", "assert f(30) == 900",
                                 "assert f(40) == 1600"], None, 1, 1))
        check("RepairLog: hard-coded answers are never learned from", len(hard.repairs) == 0)

        # ---- ZeroGPU: profile, time budget derived from the GPU call duration, memory sync
        saved = {k: os.environ.get(k) for k in ("SPACES_ZERO_GPU", "SPACE_ID", "SPACE_HARDWARE", "EIM_PROFILE", "EIM_TIME_BUDGET", "EIM_GPU_SECONDS", "EIM_GPU_SIZE")}
        try:
            for k in saved:
                os.environ.pop(k, None)
            os.environ["SPACES_ZERO_GPU"] = "1"
            hardware_profile.cache_clear()
            zp = hardware_profile()
            check("ZeroGPU is recognised: strong profile, fewer mutants, one bonus round",
                  zp.zero_gpu and zp.name == "full" and zp.mutants <= 5 and zp.bonus_iters <= 1)
            os.environ.pop("SPACES_ZERO_GPU", None)
            os.environ["SPACE_ID"] = "owner/demo"
            os.environ["SPACE_HARDWARE"] = "zerogpu"
            hardware_profile.cache_clear()
            check("ZeroGPU is recognised from Space hardware metadata before runtime allocation",
                  hardware_profile().zero_gpu and _has_accelerator())
            os.environ["SPACES_ZERO_GPU"] = "1"
            os.environ.pop("SPACE_ID", None)
            os.environ.pop("SPACE_HARDWARE", None)
            hardware_profile.cache_clear()
            os.environ.pop("EIM_GPU_SIZE", None)
            check("ZeroGPU: default budget is 65% of the 120 s large GPU call", abs(effective_time_budget(zp) - 78.0) < 1e-6)
            os.environ["EIM_GPU_SECONDS"] = "120"
            check("ZeroGPU: budget follows EIM_GPU_SECONDS", abs(effective_time_budget(zp) - 78.0) < 1e-6)
            os.environ["EIM_GPU_SECONDS"] = "180"
            check("ZeroGPU: duration and EIM budget clamp to the documented 120-second ceiling",
                  abs(effective_time_budget(zp) - 78.0) < 1e-6)
            os.environ["EIM_GPU_SIZE"] = "xlarge"
            check("ZeroGPU xlarge budget accounts for doubled quota-cost duration", abs(effective_time_budget(zp) - 39.0) < 1e-6)
            os.environ["SPACES_ZERO_GPU"] = "0"
            hardware_profile.cache_clear()
            check("a false ZeroGPU environment marker is not treated as enabled",
                  not hardware_profile().zero_gpu)
            os.environ["SPACES_ZERO_GPU"] = "1"
            hardware_profile.cache_clear()
            os.environ["EIM_TIME_BUDGET"] = "20"
            check("explicit EIM_TIME_BUDGET wins", effective_time_budget(zp) == 20.0)
            check("an explicit argument wins over everything", effective_time_budget(zp, 5) == 5.0)
            os.environ.pop("SPACES_ZERO_GPU")
            os.environ.pop("EIM_TIME_BUDGET")
            hardware_profile.cache_clear()
            check("outside ZeroGPU there is no implicit budget", effective_time_budget(hardware_profile()) == 0.0)
        finally:
            for k, v in saved.items():
                os.environ.pop(k, None)
                if v is not None:
                    os.environ[k] = v
            hardware_profile.cache_clear()

        import types
        pulled, pushed = [], []
        fake_hub = types.ModuleType("huggingface_hub")
        fake_hub.hf_hub_download = lambda **kw: pulled.append(kw["filename"])

        class _Api:
            def __init__(self, token=None): pass
            def create_repo(self, *a, **kw): pass
            def upload_file(self, **kw): pushed.append(kw["path_in_repo"])
        fake_hub.HfApi = _Api
        old_hub, old_repo, old_delay = sys.modules.get("huggingface_hub"), os.environ.get("EIM_MEMORY_REPO"), _HubSync.DELAY
        try:
            sys.modules["huggingface_hub"] = fake_hub
            os.environ["EIM_MEMORY_REPO"] = "someone/eim-memory"
            _HubSync.DELAY = 0.05
            synced = SmartMemory(path=os.path.join(tmp, "mem.json"))
            synced.add("zero|NameError", "repair")
            synced.add("zero|NameError", "rewrite")
            time.sleep(0.4)
            check("memory sync: pulled once at start, pushed (batched) after lessons",
                  pulled == ["mem.json"] and pushed == ["mem.json"])
            os.environ.pop("EIM_MEMORY_REPO")
            quiet = SmartMemory(path=os.path.join(tmp, "mem2.json"))
            quiet.add("zero|X", "repair")
            time.sleep(0.2)
            check("memory sync is off unless EIM_MEMORY_REPO is set", pulled == ["mem.json"] and pushed == ["mem.json"])
        finally:
            _HubSync.DELAY = old_delay
            if old_hub is None:
                sys.modules.pop("huggingface_hub", None)
            else:
                sys.modules["huggingface_hub"] = old_hub
            if old_repo is not None:
                os.environ["EIM_MEMORY_REPO"] = old_repo

        # ---- candidate ladder, cap, adaptive budget, time budget, result cache
        bad1, bad2 = "def add(a, b):\n    return 0", "def add(a, b):\n    return 0 * a"
        add_tests = ["assert add(1, 2) == 3", "assert add(2, 2) == 4"]
        lean0 = Profile("lean", 2, 4.0, False, 3, True, 0, 1)

        def spy_ks(engine):
            ks, original = [], engine.dcme.propose

            def spy(*args, **kwargs):
                ks.append(args[4])
                return original(*args, **kwargs)
            engine.dcme.propose = spy
            return ks

        eng = EIMPlus(_ScriptedLM([bad1, right]), eutv, fast, SmartMemory(path=""), profile=lean0)
        ks = spy_ks(eng)
        rep = list(eng.run("add", add_tests, max_iter=2, k=3))[-1]
        check("ladder: one candidate is enough when it works", ks == [1] and rep.best.result.pass_rate == 1.0)
        eng = EIMPlus(_ScriptedLM([bad1, bad2, right]), eutv, fast, SmartMemory(path=""), profile=lean0)
        ks = spy_ks(eng)
        rep = list(eng.run("add", add_tests, max_iter=2, k=3))[-1]
        check("ladder: escalates to the remaining candidates only when needed",
              ks == [1, 2] and rep.best.result.pass_rate == 1.0)
        eng = EIMPlus(_ScriptedLM([bad1, right]), eutv, fast, SmartMemory(path=""), profile=flat)
        ks = spy_ks(eng)
        list(eng.run("add", add_tests, max_iter=2, k=3))
        check("strong machines keep the full batch", ks == [3])
        rep = list(EIMPlus(_ScriptedLM([bad1]), eutv, fast, SmartMemory(path=""),
                           profile=Profile("lean", 2, 4.0, False, 2, True, 0, 0)).run("add", add_tests, max_iter=1, k=6))[-1]
        check("candidate cap is applied and explained", any("reduced from 6 to 2" in w for w in rep.warnings))

        sq_tests = ["assert f(0) == 0", "assert f(1) == 1", "assert f(2) == 4"]
        steps_lm = ["def f(x):\n    return -1", "def f(x):\n    return 0", "def f(x):\n    return x if x < 2 else 0",
                    "def f(x):\n    return x * x"]
        no_bonus = Profile("test", 4, 8.0, False, 8, False, 0, 0)
        with_bonus = Profile("test", 4, 8.0, False, 8, False, 0, 1)
        rep = list(EIMPlus(_ScriptedLM(steps_lm), eutv, fast, SmartMemory(path=""), profile=no_bonus).run(
            "square", sq_tests, max_iter=2, k=1))[-1]
        check("without the adaptive budget the loop stops at max_iter", rep.best.result.pass_rate < 1.0)
        rep = list(EIMPlus(_ScriptedLM(steps_lm), eutv, fast, SmartMemory(path=""), profile=with_bonus).run(
            "square", sq_tests, max_iter=2, k=1))[-1]
        check("adaptive budget: an extra round while the code is still improving",
              rep.best.result.pass_rate == 1.0 and any("extended" in w for w in rep.warnings))

        rep = list(EIMPlus(_ScriptedLM([bad1, right]), eutv, fast, SmartMemory(path=""), profile=flat).run(
            "add", add_tests, max_iter=3, k=1, time_budget=1e-9))[-1]
        check("time budget: returns the best version found so far", rep.iterations == 0
              and any("time budget" in w for w in rep.warnings))

        lm = _ScriptedLM([bad1, right])
        eng = EIMPlus(lm, eutv, fast, SmartMemory(path=""), profile=flat)
        first_events = list(eng.run("add", add_tests, max_iter=2, k=1))
        used = lm.i
        second_events = list(eng.run("add", add_tests, max_iter=2, k=1))
        check("a fully successful run is replayed from the cache without calling the model",
              lm.i == used and second_events == first_events)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"\nAll {checks} checks passed.")
    return 0


def _bomb_docx(tmp: str) -> str:
    path = os.path.join(tmp, "bomb.docx")
    xml = ('<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><w:document xmlns:w="http://schemas.openxmlformats.org/'
           'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&a;</w:t></w:r></w:p></w:body></w:document>')
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", xml)
    return path


if __name__ == "__main__" and "--selftest" in sys.argv:
    raise SystemExit(run_selftest())

if __name__ == "__main__" and "--export-sft" in sys.argv:
    # python eim_plus.py --export-sft train.jsonl   -> verified repairs as chat-format SFT data (LoRA/QLoRA)
    target = sys.argv[sys.argv.index("--export-sft") + 1] if len(sys.argv) > sys.argv.index("--export-sft") + 1 else "sft.jsonl"
    repair_log = RepairLog.for_config(CFG)
    if repair_log is None:
        raise SystemExit("no repair log: set a memory path in the config or EIM_REPAIR_LOG_PATH")
    print(f"{export_sft(repair_log, target)} verified examples written to {target}")
    raise SystemExit(0)
